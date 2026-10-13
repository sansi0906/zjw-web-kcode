import ast
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

test_dir = Path('tests')
test_files = sorted(test_dir.glob('test_*.py'))

# Mapping from fixture/file names to Chinese page/module names
PAGE_NAME_MAP = {
    'achievement_transformation': '成果转化',
    'industry_association': '产业协会',
    'industry_policy': '产业政策',
    'intelligent_construction': '智能建造',
    'cloud_service': '云服务',
    'all_products': '全部商品',
    'industrial_service': '产业化服务',
    'industry_demand': '产业需求',
    'industry_news': '产业资讯',
    'key_projects': '重点项目',
    'supply_hall': '供应大厅',
    'supply_chain_map': '供应链地图',
    'homepage': '首页',
}

MODULE_NAME_MAP = {
    'achievement_transformation': '成果转化',
    'industry_association': '产业协会',
    'industry_policy': '产业政策',
    'intelligent_construction': '智能建造',
    'cloud_service': '云服务',
    'all_products': '全部商品',
    'industrial_service': '产业化服务',
    'industry_demand': '产业需求',
    'industry_news': '产业资讯',
    'key_projects': '重点项目',
    'supply_hall': '供应大厅',
    'supply_chain_map': '供应链地图',
    'homepage': '首页',
}

SEVERITY_MAP = {
    'BLOCKER': '阻塞',
    'CRITICAL': '严重',
    'NORMAL': '一般',
    'MINOR': '轻微',
    'TRIVIAL': '提示',
}

rows = []
idx = 1

for filepath in test_files:
    rel = filepath.relative_to(test_dir).as_posix()
    stem = filepath.stem
    try:
        source = filepath.read_text(encoding='utf-8')
    except Exception:
        continue
    try:
        tree = ast.parse(source)
    except SyntaxError:
        continue

    for node in tree.body:
        if isinstance(node, ast.ClassDef):
            feature = ''
            page_name = ''
            fixture_name = ''
            for decorator in node.decorator_list:
                if isinstance(decorator, ast.Call) and getattr(decorator.func, 'attr', '') == 'feature':
                    try:
                        feature = decorator.args[0].value
                    except Exception:
                        pass

            for item in node.body:
                if isinstance(item, ast.Assign):
                    target = item.targets[0]
                    if isinstance(target, ast.Name):
                        name = target.id
                        if name == 'page_name':
                            try:
                                page_name = item.value.value
                            except Exception:
                                pass
                        elif name == 'page_fixture_name':
                            try:
                                fixture_name = item.value.value
                            except Exception:
                                pass
                if isinstance(item, ast.FunctionDef) and item.name.startswith('test_'):
                    method_doc = ast.get_docstring(item) or ''
                    title = ''
                    story = ''
                    severity = ''
                    for decorator in item.decorator_list:
                        if isinstance(decorator, ast.Call) and getattr(decorator.func, 'attr', '') == 'title':
                            try:
                                title = decorator.args[0].value
                            except Exception:
                                pass
                        elif isinstance(decorator, ast.Call) and getattr(decorator.func, 'attr', '') == 'story':
                            try:
                                story = decorator.args[0].value
                            except Exception:
                                pass
                        elif isinstance(decorator, ast.Call) and getattr(decorator.func, 'attr', '') == 'severity':
                            try:
                                sev_arg = decorator.args[0]
                                if isinstance(sev_arg, ast.Attribute):
                                    severity = sev_arg.attr
                            except Exception:
                                pass

                    severity = SEVERITY_MAP.get(severity.upper(), severity) if severity else '一般'

                    # Determine module and page
                    module = page_name or MODULE_NAME_MAP.get(stem.replace('test_', ''), feature or rel)
                    page = PAGE_NAME_MAP.get(fixture_name.replace('_page', ''), page_name or module)

                    rows.append({
                        '用例编号': f'TC-{idx:03d}',
                        '模块': module,
                        '页面': page,
                        '用例名称': item.name,
                        '用例标题': title or (method_doc.split('\n')[0] if method_doc else item.name),
                        'Feature': feature,
                        'Story': story,
                        '严重程度': severity,
                        '测试类型': '功能',
                        '前置条件': '打开对应页面',
                        '测试步骤': method_doc.replace('\n', '; ') if method_doc else '',
                        '预期结果': '',
                        '测试状态': '',
                        '备注': '',
                    })
                    idx += 1

print(f'Collected {len(rows)} test cases')

# Write to Excel
wb = Workbook()
ws = wb.active
ws.title = '测试用例'

headers = [
    '用例编号', '模块', '页面', '用例名称', '用例标题', 'Feature', 'Story',
    '严重程度', '测试类型', '前置条件', '测试步骤', '预期结果', '测试状态', '备注'
]

# Header style
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font = Font(color='FFFFFF', bold=True)
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_align

# Data
for row_idx, row in enumerate(rows, 2):
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=row.get(header, ''))
        cell.alignment = Alignment(vertical='center', wrap_text=True)

# Column widths
col_widths = [10, 18, 12, 25, 30, 20, 20, 10, 10, 15, 40, 30, 10, 15]
for i, width in enumerate(col_widths, 1):
    ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

# Freeze header
ws.freeze_panes = 'A2'

output_path = 'reports/测试用例清单.xlsx'
wb.save(output_path)
print(f'Saved to {output_path}')
print(f'Total test cases: {len(rows)}')

# Verify Chinese content
print('\nSample rows:')
for row in rows[:5]:
    print(row['用例编号'], row['模块'], row['页面'], row['用例名称'], row['用例标题'][:20])